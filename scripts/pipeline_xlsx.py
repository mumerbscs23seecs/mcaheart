"""
Excel round-trip for the pipeline's Submissions, Ongoing and Conferences pages.

  python3 scripts/pipeline_xlsx.py export OUT_DIR
      Writes Submissions.xlsx, Ongoing.xlsx, Conferences.xlsx - the same rows
      the site shows, in the same order.

  python3 scripts/pipeline_xlsx.py import FILE.xlsx [--apply]
      Reads one of those files back (edited) and applies the differences.
      Dry run by default: prints every change and every skipped row first.

Rows are matched by Ref (MCA-xxxx), never by title. White columns are
editable, grey ones are reference only and ignored on import. Status/stage
changes and outcomes go through the same database functions the website's
own buttons call (no emails are sent); plain metadata (short title, order)
is written directly with the service key, as the earlier sync scripts do.
Nothing is ever deleted: a row missing from the file is just reported.
"""
import json
import os
import sys
import urllib.request
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ADMIN_EMAIL = 'admin@mcaheart.com'
STUDY_TYPES = ['DMC PCI', 'TriNetX', 'SRMA', 'NIS', 'NRD', 'CDC', 'Other']


def load_env():
    env = {}
    with open(os.path.join(ROOT, '.env')) as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith('#') and '=' in line:
                k, v = line.split('=', 1)
                env[k.strip()] = v.strip().strip('"').strip("'")
    return env


ENV = load_env()
URL = ENV['SUPABASE_URL'].rstrip('/')
SERVICE = ENV['SUPABASE_SERVICE_ROLE_KEY']
ANON = ENV['SUPABASE_ANON_KEY']


def http(method, path, body=None, key=SERVICE, token=None, extra=None):
    headers = {'apikey': key, 'Authorization': f'Bearer {token or key}', 'Content-Type': 'application/json'}
    headers.update(extra or {})
    req = urllib.request.Request(URL + path, method=method, headers=headers,
                                 data=json.dumps(body).encode() if body is not None else None)
    try:
        with urllib.request.urlopen(req) as r:
            raw = r.read()
            return json.loads(raw) if raw else None
    except urllib.error.HTTPError as e:
        raise RuntimeError(f'{method} {path}: {e.read().decode()}') from None


def get(table, query):
    return http('GET', f'/rest/v1/{table}?{query}')


def patch(table, match, body):
    return http('PATCH', f'/rest/v1/{table}?{match}', body, extra={'Prefer': 'return=minimal'})


_token = None


def rpc(fn, args):
    """Calls an RPC as the real admin user - the functions check me()/auth.uid()."""
    global _token
    if _token is None:
        link = http('POST', '/auth/v1/admin/generate_link', {'type': 'magiclink', 'email': ADMIN_EMAIL})
        hashed = link.get('hashed_token') or link['properties']['hashed_token']
        sess = http('POST', '/auth/v1/verify', {'type': 'email', 'token_hash': hashed}, key=ANON)
        _token = sess['access_token']
    return http('POST', f'/rest/v1/rpc/{fn}', args, key=ANON, token=_token)


# --------------------------------------------------------------------- data
def stages():
    rows = get('stages', 'select=phase,code,label,sort&order=sort')
    out = {}
    for r in rows:
        if r['phase'] == 'conference' and r['code'] == 'abstract_wip':
            continue  # retired - never offered
        out.setdefault(r['phase'], []).append(r)
    return out


def people():
    return get('people', 'select=id,full_name&active=eq.true&order=full_name')


PL_COLS = ('id,ref,title,full_title,phase,stage,stage_label,study_type,lead_id,analyst_id,colead_id,'
           'presenter_id,corresponding_id,lead_name,analyst_name,current_journal,current_conference,'
           'journals_tried,open_journal_attempt_id,parked,submission_order,stage_changed_at')


def submissions_rows():
    return get('project_list', f'select={PL_COLS}&phase=eq.journal&stage=neq.published&archived_at=is.null'
               '&hidden_at=is.null&order=submission_order.asc.nullslast,stage_changed_at.desc')


def ongoing_rows():
    return get('project_list', f'select={PL_COLS}&phase=eq.manuscript&archived_at=is.null'
               '&hidden_at=is.null&order=submission_order.asc.nullslast,stage_changed_at.desc')


def journal_history(ids):
    if not ids:
        return {}
    rows = get('journal_attempts', 'select=project_id,journal,outcome,submitted_on&project_id=in.('
               + ','.join(ids) + ')&order=submitted_on.asc,id.asc')
    out = {}
    for r in rows:
        out.setdefault(r['project_id'], []).append(f"{r['journal']} ({r['outcome'] or 'open'})")
    return out


def conferences_with_rows():
    confs = get('conferences', 'select=id,name,abstract_deadline&order=abstract_deadline.desc')
    attempts = get('conference_attempts', 'select=id,project_id,conference_id,outcome')
    ids = sorted({a['project_id'] for a in attempts})
    projs = {p['id']: p for p in get('project_list', f'select={PL_COLS},hidden_at&id=in.(' + ','.join(ids) + ')')} if ids else {}
    group = {'accepted': 1, 'rejected': 2, 'withdrawn': 2}
    out = []
    for c in confs:
        rows = []
        for a in attempts:
            p = projs.get(a['project_id'])
            if a['conference_id'] != c['id'] or not p or p.get('hidden_at'):
                continue
            rows.append((a, p))
        rows.sort(key=lambda r: (group.get(r[0]['outcome'], 3) if r[0]['outcome'] else 0, r[1]['title'].lower()))
        out.append((c, rows))
    return out


# ------------------------------------------------------------------- styling
FONT = Font(name='Arial', size=10)
BOLD = Font(name='Arial', size=10, bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='A51C30')  # site crimson
GREY = PatternFill('solid', fgColor='EDEDED')
GREY_FONT = Font(name='Arial', size=10, color='555555')
THIN = Side(style='thin', color='D0D0D0')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_lists(wb, st, ppl):
    ws = wb.create_sheet('Lists')
    cols = {'Type': STUDY_TYPES, 'People': [p['full_name'] for p in ppl]}
    for phase in ('journal', 'manuscript', 'conference'):
        cols[f'Status {phase}'] = [s['label'] for s in st.get(phase, [])]
    refs = {}
    for i, (name, values) in enumerate(cols.items(), start=1):
        letter = get_column_letter(i)
        ws.cell(row=1, column=i, value=name).font = Font(name='Arial', size=10, bold=True)
        for j, v in enumerate(values, start=2):
            ws.cell(row=j, column=i, value=v).font = FONT
        refs[name] = f"Lists!${letter}$2:${letter}${len(values) + 1}"
    ws.sheet_state = 'hidden'
    return refs


def write_table(ws, cols, rows, lists):
    """cols: [(header, width, editable, list_name_or_None)]"""
    for i, (h, w, editable, _) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.border = BOLD, HEAD_FILL, BOX
        c.alignment = Alignment(vertical='center', wrap_text=True)
        if not editable:
            c.comment = Comment('Reference only - edits here are ignored on import.', 'MCA pipeline')
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    for r, values in enumerate(rows, start=2):
        for i, (v, (_, _, editable, _)) in enumerate(zip(values, cols), start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font = FONT if editable else GREY_FONT
            c.border = BOX
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if not editable:
                c.fill = GREY
    last = max(len(rows) + 1, 2) + 200  # room to add rows below
    for i, (_, _, editable, list_name) in enumerate(cols, start=1):
        if editable and list_name:
            dv = DataValidation(type='list', formula1='=' + lists[list_name], allow_blank=True,
                                showErrorMessage=True, errorTitle='Not in the list',
                                error='Pick a value from the dropdown.')
            letter = get_column_letter(i)
            dv.add(f'{letter}2:{letter}{last}')
            ws.add_data_validation(dv)
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(rows) + 1}'


def write_howto(wb, page, editable_notes):
    ws = wb.active
    ws.title = 'How to edit'
    lines = [
        (f'MCA pipeline - {page}', Font(name='Arial', size=14, bold=True)),
        (f'Exported {date.today().isoformat()} from the live pipeline.', FONT),
        ('', FONT),
        ('White columns are editable. Grey columns are for reference and are ignored when the file comes back.', FONT),
        ('Ref (MCA-xxxx) is how each row is matched back - never change or reuse it.', FONT),
        ('Type, Lead, Analyst and Status have dropdowns - pick from them so names and statuses match exactly.', FONT),
        ('Deleting a row does NOT delete the paper; it is just reported. To hide a paper, use Hide on its page.', FONT),
        ('', FONT),
    ] + [(n, FONT) for n in editable_notes]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = font
    ws.column_dimensions['A'].width = 110


def export(out_dir):
    os.makedirs(out_dir, exist_ok=True)
    st, ppl = stages(), people()
    names = {p['id']: p['full_name'] for p in ppl}

    # ---- Submissions
    subs = submissions_rows()
    hist = journal_history([p['id'] for p in subs])
    wb = Workbook()
    write_howto(wb, 'Submissions', [
        'Row order = order on the Submissions page (drag/cut rows to reorder).',
        'Journal: type a new journal name to record a resubmission there (the current one must already be decided, e.g. Rejected).',
        'Status Accepted / Rejected records the outcome on the current journal.',
    ])
    lists = write_lists(wb, st, ppl)
    ws = wb.create_sheet('Submissions', 1)
    write_table(ws, [
        ('Ref', 11, False, None), ('Submission title', 60, True, None), ('Abbreviation', 28, True, None),
        ('Type', 11, True, 'Type'), ('Lead', 20, True, 'People'), ('Analyst', 20, True, 'People'),
        ('Journal', 34, True, None), ('Status', 24, True, 'Status journal'), ('Journals tried', 48, False, None),
    ], [[p['ref'], p['full_title'] or '', p['title'], p['study_type'], names.get(p['lead_id'], p['lead_name']),
         names.get(p['analyst_id'], p['analyst_name']), p['current_journal'], p['stage_label'],
         '  >  '.join(hist.get(p['id'], []))] for p in subs], lists)
    wb.save(os.path.join(out_dir, 'Submissions.xlsx'))

    # ---- Ongoing
    ong = ongoing_rows()
    wb = Workbook()
    write_howto(wb, 'Ongoing', ['Row order = order on the Ongoing page (drag/cut rows to reorder).'])
    lists = write_lists(wb, st, ppl)
    ws = wb.create_sheet('Ongoing', 1)
    write_table(ws, [
        ('Ref', 11, False, None), ('Submission title', 60, True, None), ('Abbreviation', 28, True, None),
        ('Type', 11, True, 'Type'), ('Lead', 20, True, 'People'), ('Analyst', 20, True, 'People'),
        ('Conference', 14, False, None), ('Status', 40, True, 'Status manuscript'), ('Parked', 9, False, None),
    ], [[p['ref'], p['full_title'] or '', p['title'], p['study_type'], names.get(p['lead_id'], p['lead_name']),
         names.get(p['analyst_id'], p['analyst_name']), p['current_conference'], p['stage_label'],
         'yes' if p['parked'] else ''] for p in ong], lists)
    wb.save(os.path.join(out_dir, 'Ongoing.xlsx'))

    # ---- Conferences: one tab per conference, newest first
    wb = Workbook()
    write_howto(wb, 'Conferences', [
        'One tab per conference, same rows as that conference\'s page.',
        'Status only applies to papers still in that conference and undecided; for the rest it shows the outcome (grey).',
        'New paper: add a row with no Ref - Abbreviation and Type are required; it is created straight into that conference.',
    ])
    lists = write_lists(wb, st, ppl)
    for idx, (c, rows) in enumerate(conferences_with_rows(), start=1):
        ws = wb.create_sheet(c['name'][:31], idx)
        data = []
        for a, p in rows:
            live = not a['outcome'] and p['phase'] == 'conference'
            data.append([p['ref'], p['title'], p['study_type'], names.get(p['lead_id'], p['lead_name']),
                         names.get(p['analyst_id'], p['analyst_name']),
                         p['stage_label'] if live else (a['outcome'] or 'pending').capitalize(),
                         {'conference': 'Conferences', 'manuscript': 'Ongoing', 'journal': 'Submissions',
                          'idea': 'Ideas'}.get(p['phase'], p['phase'])])
        write_table(ws, [
            ('Ref', 11, False, None), ('Abbreviation', 44, True, None), ('Type', 11, True, 'Type'),
            ('Lead', 20, True, 'People'), ('Analyst', 20, True, 'People'),
            ('Status', 24, True, 'Status conference'), ('Now in', 13, False, None),
        ], data, lists)
        # Decided / moved-on rows: their Status is an outcome, not editable.
        for r, (a, p) in enumerate(rows, start=2):
            if a['outcome'] or p['phase'] != 'conference':
                cell = ws.cell(row=r, column=6)
                cell.fill, cell.font = GREY, GREY_FONT
    wb.save(os.path.join(out_dir, 'Conferences.xlsx'))
    print(f'Wrote Submissions.xlsx ({len(subs)} rows), Ongoing.xlsx ({len(ong)} rows), Conferences.xlsx to {out_dir}')


# -------------------------------------------------------------------- import
def norm(s):
    return ' '.join(str(s or '').split()).lower()


def read_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or '').strip() for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not any(v not in (None, '') for v in r):
            continue
        out.append({h: (str(v).strip() if v is not None else '') for h, v in zip(head, r)})
    return out


class Plan:
    def __init__(self, apply):
        self.apply, self.changes, self.flags = apply, [], []

    def do(self, label, fn):
        self.changes.append(label)
        if self.apply:
            # One bad row (e.g. a title another paper already uses) must not stop
            # the rest of the file half-applied - report it and keep going.
            try:
                fn()
            except RuntimeError as e:
                self.changes.pop()
                self.flag(f'{label} - FAILED: {str(e)[-160:]}')

    def quiet(self, fn):
        """A change counted in one summary line instead of listed (row reordering)."""
        self.reorders = getattr(self, 'reorders', 0) + 1
        if self.apply:
            fn()

    def flag(self, label):
        self.flags.append(label)


def common_fields(plan, p, row, names_to_id, title_col):
    tag = f"{p['ref']} {p['title']}"
    if title_col in row and row[title_col] and row[title_col] != p['title']:
        plan.do(f'{tag}: abbreviation -> "{row[title_col]}"',
                lambda: patch('projects', f"id=eq.{p['id']}", {'title': row[title_col]}))
    full_col = 'Submission title' if 'Submission title' in row else 'Full title' if 'Full title' in row else None
    # Blank means "no change" (same as the site's full-title box) - clear one on the paper page.
    if full_col and row[full_col] and row[full_col] != (p['full_title'] or ''):
        plan.do(f'{tag}: full title -> "{row[full_col]}"',
                lambda: rpc('set_full_title', {'p_project': p['id'], 'p_full_title': row[full_col]}))
    if row.get('Type') and row['Type'] != p['study_type']:
        if row['Type'] in STUDY_TYPES:
            plan.do(f"{tag}: type {p['study_type']} -> {row['Type']}",
                    lambda: rpc('set_study_type', {'p_project': p['id'], 'p_study_type': row['Type']}))
        else:
            plan.flag(f'{tag}: unknown type "{row["Type"]}"')
    new = {'lead_id': p['lead_id'], 'analyst_id': p['analyst_id']}
    for col, key in (('Lead', 'lead_id'), ('Analyst', 'analyst_id')):
        if col not in row:
            continue
        want = row[col]
        cur = names_to_id.get('__id__' + str(p[key]))
        if norm(want) == norm(cur):
            continue
        if not want:
            new[key] = None
        elif norm(want) in names_to_id:
            new[key] = names_to_id[norm(want)]
        else:
            plan.flag(f'{tag}: {col.lower()} "{want}" is not a person in the system - left as "{cur or "-"}"')
    if new['lead_id'] != p['lead_id'] or new['analyst_id'] != p['analyst_id']:
        plan.do(f"{tag}: lead/analyst -> {row.get('Lead') or '-'} / {row.get('Analyst') or '-'}",
                lambda: rpc('assign_people', {
                    'p_project': p['id'], 'p_lead_id': new['lead_id'], 'p_colead_id': p['colead_id'],
                    'p_analyst_id': new['analyst_id'], 'p_presenter_id': p['presenter_id'],
                    'p_corresponding_id': p['corresponding_id']}))


def import_file(path, apply):
    wb = load_workbook(path, data_only=True)
    st, ppl = stages(), people()
    counts = {}
    for x in ppl:
        counts[norm(x['full_name'])] = counts.get(norm(x['full_name']), 0) + 1
    # A name shared by two people can't be matched safely - treated as unknown.
    names_to_id = {norm(x['full_name']): x['id'] for x in ppl if counts[norm(x['full_name'])] == 1}
    names_to_id.update({'__id__' + p['id']: p['full_name'] for p in ppl})
    by_label = {ph: {norm(s['label']): s['code'] for s in ss} for ph, ss in st.items()}
    plan = Plan(apply)
    kind = 'Submissions' if 'Submissions' in wb.sheetnames else 'Ongoing' if 'Ongoing' in wb.sheetnames else 'Conferences'

    if kind in ('Submissions', 'Ongoing'):
        live = submissions_rows() if kind == 'Submissions' else ongoing_rows()
        by_ref = {p['ref']: p for p in live}
        rows = read_sheet(wb[kind])
        # Only rewrite order when the sequence of papers actually changed -
        # many rows have no stored order yet and already display correctly.
        reordered = [r.get('Ref') for r in rows if r.get('Ref') in by_ref] != [p['ref'] for p in live if p['ref'] in {r.get('Ref') for r in rows}]
        seen = set()
        for i, row in enumerate(rows):
            p = by_ref.get(row.get('Ref', ''))
            if not p:
                plan.flag(f"row {i + 2}: Ref \"{row.get('Ref', '')}\" is not on the {kind} page - skipped")
                continue
            seen.add(p['ref'])
            tag = f"{p['ref']} {p['title']}"
            common_fields(plan, p, row, names_to_id, 'Abbreviation')
            if reordered and p['submission_order'] != i:
                plan.quiet(lambda p=p, i=i: patch('projects', f"id=eq.{p['id']}", {'submission_order': i}))
            phase = 'journal' if kind == 'Submissions' else 'manuscript'
            want = by_label[phase].get(norm(row.get('Status')))
            if row.get('Status') and not want:
                plan.flag(f'{tag}: unknown status "{row["Status"]}"')
                continue
            journal = row.get('Journal', '')
            if kind == 'Submissions' and journal and norm(journal) != norm(p['current_journal']):
                rejected = want in ('rejected_comments', 'rejected_no_comments')
                if p['open_journal_attempt_id'] and not rejected:
                    plan.flag(f'{tag}: journal changed to "{journal}" but "{p["current_journal"]}" is still open - '
                              'set Status to Rejected to record that and resubmit, or fix the name on the paper page')
                    continue
                if p['open_journal_attempt_id']:
                    # Rejected at the open journal AND resubmitted: one call closes
                    # the old attempt and opens the new one at Submitted.
                    plan.do(f'{tag}: rejected at "{p["current_journal"]}", resubmitted to "{journal}"',
                            lambda p=p, j=journal, w=want: rpc('record_journal_outcome', {
                                'p_attempt': p['open_journal_attempt_id'], 'p_outcome': 'rejected',
                                'p_next_journal': j, 'p_reject_stage': w}))
                else:
                    plan.do(f'{tag}: resubmitted to "{journal}"',
                            lambda p=p, j=journal: rpc('set_submitted_journal', {'p_project': p['id'], 'p_journal': j}))
                    # A status left on the old outcome just means "not updated yet";
                    # a live stage (Under review...) is applied on top of Submitted.
                    if want and want not in ('submitted', 'accepted', 'rejected_comments', 'rejected_no_comments', 'published'):
                        plan.do(f'{tag}: status -> {row["Status"]}',
                                lambda p=p, w=want: rpc('advance_stage', {'p_project': p['id'], 'p_stage': w}))
                continue
            if not want or want == p['stage']:
                continue
            outcome = {'accepted': ('accepted', None), 'rejected_comments': ('rejected', 'rejected_comments'),
                       'rejected_no_comments': ('rejected', 'rejected_no_comments')}.get(want) if kind == 'Submissions' else None
            if outcome:
                if not p['open_journal_attempt_id']:
                    plan.flag(f'{tag}: {row["Status"]} needs an open journal submission - none is open')
                    continue
                plan.do(f'{tag}: outcome {row["Status"]} at {p["current_journal"]}',
                        lambda p=p, o=outcome: rpc('record_journal_outcome', {
                            'p_attempt': p['open_journal_attempt_id'], 'p_outcome': o[0],
                            'p_next_journal': None, 'p_reject_stage': o[1]}))
            elif kind == 'Submissions' and want == 'submitted':
                plan.do(f'{tag}: status -> {row["Status"]}',
                        lambda p=p: rpc('set_submitted_journal', {'p_project': p['id'], 'p_journal': p['current_journal']}))
            else:
                plan.do(f'{tag}: status {p["stage_label"]} -> {row["Status"]}',
                        lambda p=p, w=want: rpc('advance_stage', {'p_project': p['id'], 'p_stage': w}))
        for ref in sorted(set(by_ref) - seen):
            plan.flag(f'{ref} {by_ref[ref]["title"]}: not in the file - left unchanged')

    else:  # Conferences
        confs = {c['name'][:31]: (c, rows) for c, rows in conferences_with_rows()}
        for name in wb.sheetnames:
            if name not in confs:
                continue
            c, live_rows = confs[name]
            by_ref = {p['ref']: (a, p) for a, p in live_rows}
            for i, row in enumerate(read_sheet(wb[name])):
                ref = row.get('Ref', '')
                if not ref:
                    title, typ = row.get('Abbreviation', ''), row.get('Type', '')
                    if not title or typ not in STUDY_TYPES:
                        plan.flag(f'{name} row {i + 2}: new paper needs an Abbreviation and a Type - skipped')
                        continue
                    lead, analyst = names_to_id.get(norm(row.get('Lead'))), names_to_id.get(norm(row.get('Analyst')))
                    want = by_label['conference'].get(norm(row.get('Status')))

                    def create(title=title, typ=typ, lead=lead, analyst=analyst, want=want, c=c):
                        pr = rpc('create_project_in_conference', {
                            'p_title': title, 'p_study_type': typ, 'p_conference_id': c['id'],
                            'p_lead_id': lead, 'p_analyst_id': analyst})
                        if want and want not in ('no_progress', 'accepted', 'rejected'):
                            rpc('advance_stage', {'p_project': pr['id'], 'p_stage': want})
                    plan.do(f'{name}: NEW paper "{title}" ({typ}{", " + row["Status"] if want else ""})', create)
                    continue
                if ref not in by_ref:
                    plan.flag(f'{name} row {i + 2}: Ref "{ref}" is not in this conference - skipped')
                    continue
                a, p = by_ref[ref]
                tag = f"{name} {p['ref']} {p['title']}"
                common_fields(plan, p, row, names_to_id, 'Abbreviation')
                if a['outcome'] or p['phase'] != 'conference':
                    continue  # status is an outcome here, read-only
                want = by_label['conference'].get(norm(row.get('Status')))
                if row.get('Status') and not want:
                    plan.flag(f'{tag}: unknown status "{row["Status"]}"')
                elif want and want != p['stage']:
                    if want in ('accepted', 'rejected'):
                        plan.do(f'{tag}: abstract {want}',
                                lambda a=a, w=want: rpc('record_abstract_outcome', {
                                    'p_attempt': a['id'], 'p_outcome': w, 'p_presenter': None}))
                    else:
                        plan.do(f'{tag}: status {p["stage_label"]} -> {row["Status"]}',
                                lambda p=p, w=want: rpc('advance_stage', {'p_project': p['id'], 'p_stage': w}))

    if getattr(plan, 'reorders', 0):
        plan.changes.append(f'row order changed - {plan.reorders} paper(s) renumbered to match the file')
    print(f'=== {kind}: {len(plan.changes)} change(s) ===')
    for c in plan.changes:
        print(' ' + c)
    print(f'\n=== For you to check ({len(plan.flags)}) - nothing changed for these ===')
    for f in plan.flags:
        print(' ' + f)
    print('\nApplied.' if apply else '\n(dry run - nothing written; re-run with --apply)')


if __name__ == '__main__':
    if len(sys.argv) >= 3 and sys.argv[1] == 'export':
        export(sys.argv[2])
    elif len(sys.argv) >= 3 and sys.argv[1] == 'import':
        import_file(sys.argv[2], '--apply' in sys.argv)
    else:
        print(__doc__)
        sys.exit(1)
